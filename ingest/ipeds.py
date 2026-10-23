# IPEDS complete data files for the 18 peers. URLs come from config/ipeds_files.csv, copied
# from the Data Center's year pages: since late 2025 older and newer years live in different
# folders, and the old folder keeps stale copies of the newer files, so no URL is built from a
# pattern.

import argparse
import csv
import hashlib
import io
import logging
import sys
import zipfile
from datetime import UTC, datetime

import openpyxl
import pandas as pd

from ingest.http import fetch, new_session, require_ok
from ingest.paths import CONFIG, EXTRACTS, RAW

log = logging.getLogger("ipeds")

RAW_IPEDS = RAW / "ipeds"
EXTRACTS_IPEDS = EXTRACTS / "ipeds"
MANIFEST = EXTRACTS / "manifest.csv"
MANIFEST_COLUMNS = [
    "source",
    "url",
    "retrieved_utc",
    "last_modified",
    "etag",
    "bytes",
    "sha256",
    "file_used",
]
DICTIONARY_TITLES = EXTRACTS_IPEDS / "dictionary_titles.csv"


def read_peer_unitids():
    with (CONFIG / "peers.csv").open(newline="") as f:
        unitids = {row["unitid"] for row in csv.DictReader(f)}
    assert len(unitids) == 18, f"config/peers.csv lists {len(unitids)} schools, expected 18"
    return unitids


def read_files():
    with (CONFIG / "ipeds_files.csv").open(newline="") as f:
        files = list(csv.DictReader(f))
    names = [row["file"] for row in files]
    assert len(names) == len(set(names)), "config/ipeds_files.csv lists a file twice"
    return files


def download(session, url, dest):
    response = fetch(session, url)
    content = require_ok(response)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(content)
    log.info("%s: %d bytes", dest.name, len(content))
    return {
        "source": f"ipeds/{dest.name}",
        "url": url,
        "retrieved_utc": datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "last_modified": response.headers.get("Last-Modified", ""),
        "etag": response.headers.get("ETag", ""),
        "bytes": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
    }


def update_manifest(new_rows):
    rows = {}
    if MANIFEST.exists():
        with MANIFEST.open(newline="") as f:
            rows = {row["source"]: row for row in csv.DictReader(f)}
    rows.update({row["source"]: row for row in new_rows})
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    with MANIFEST.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=MANIFEST_COLUMNS)
        writer.writeheader()
        writer.writerows(rows[source] for source in sorted(rows))


def sheet_rows(workbook, name):
    # Sheet names and header names change case between years (Varlist and varlist,
    # varNumber and varnumber), so both are matched in lower case.
    sheets = {sheet.lower(): sheet for sheet in workbook.sheetnames}
    if name not in sheets:
        return []
    rows = list(workbook[sheets[name]].iter_rows(values_only=True))
    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    return [dict(zip(header, row, strict=False)) for row in rows[1:] if any(row)]


def read_dictionary(zip_path):
    with zipfile.ZipFile(zip_path) as z:
        workbooks = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        assert len(workbooks) == 1, f"{zip_path.name}: expected one .xlsx, found {workbooks}"
        workbook = openpyxl.load_workbook(io.BytesIO(z.read(workbooks[0])), read_only=True)
    # Matched by prefix: GR2017_PELL_SSL spells the sheet "Introdution".
    intros = [sheet for sheet in workbook.sheetnames if sheet.lower().startswith("intro")]
    assert len(intros) == 1, f"{zip_path.name}: expected one introduction sheet, found {intros}"
    intro = [row[0] for row in workbook[intros[0]].iter_rows(values_only=True)]
    title, release = intro[0], intro[1]
    assert "documentation" in str(title).lower(), f"{zip_path.name}: unexpected title {title!r}"
    assert str(release).startswith("("), f"{zip_path.name}: unexpected release line {release!r}"
    descriptions = {
        row["varname"]: row.get("longdescription") for row in sheet_rows(workbook, "description")
    }
    variables = {
        row["varname"]: {"title": row["vartitle"], "description": descriptions.get(row["varname"])}
        for row in sheet_rows(workbook, "varlist")
    }
    codes = {}
    for row in sheet_rows(workbook, "frequencies"):
        codes.setdefault(row["varname"], {})[str(row["codevalue"])] = row["valuelabel"]
    return {
        "workbook": workbooks[0],
        "title": title,
        "release": release,
        "variables": variables,
        "codes": codes,
    }


def download_dictionaries():
    session = new_session()
    manifest_rows, titles = [], []
    for row in read_files():
        dest = RAW_IPEDS / f"{row['file']}_Dict.zip"
        manifest_row = download(session, row["dict_url"], dest)
        dictionary = read_dictionary(dest)
        manifest_rows.append(manifest_row | {"file_used": dictionary["workbook"]})
        titles.append(
            {"file": row["file"], "title": dictionary["title"], "release": dictionary["release"]}
        )
    update_manifest(manifest_rows)
    EXTRACTS_IPEDS.mkdir(parents=True, exist_ok=True)
    with DICTIONARY_TITLES.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["file", "title", "release"])
        writer.writeheader()
        writer.writerows(titles)
    log.info("downloaded %d dictionaries", len(titles))


def choose_csv(names):
    # Once NCES revises a file, the zip holds the original CSV and a revised one whose name
    # ends in _rv or _RV. Any other layout is unexpected and stops the run.
    csvs = [n for n in names if n.lower().endswith(".csv")]
    revised = [n for n in csvs if n.lower().endswith("_rv.csv")]
    if revised:
        assert len(revised) == 1 and len(csvs) == 2, f"unexpected CSVs in zip: {csvs}"
        return revised[0]
    assert len(csvs) == 1, f"unexpected CSVs in zip: {csvs}"
    return csvs[0]


def decode(raw, label):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{label} is neither utf-8 nor cp1252")


def extract(file, zip_path, unitids):
    with zipfile.ZipFile(zip_path) as z:
        name = choose_csv(z.namelist())
        text, encoding = decode(z.read(name), name)
    # Every column stays text so nothing is coerced on the way in: CIPCODE keeps its
    # NN.NNNN form and blanks stay blank. Typing happens in dbt staging.
    source = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False)
    assert "UNITID" in source.columns, f"{name}: no UNITID column in {list(source.columns)[:5]}"
    kept = source[source["UNITID"].str.strip().isin(unitids)]
    missing = sorted(unitids - set(kept["UNITID"].str.strip()))
    EXTRACTS_IPEDS.mkdir(parents=True, exist_ok=True)
    kept.to_parquet(EXTRACTS_IPEDS / f"{file}.parquet", index=False)
    log.info(
        "%s: used %s (%s), kept %d of %d rows, %d of 18 schools",
        file,
        name,
        encoding,
        len(kept),
        len(source),
        18 - len(missing),
    )
    if missing:
        log.warning("%s: no rows for UNITIDs %s", file, missing)
    return name


def download_data():
    session = new_session()
    unitids = read_peer_unitids()
    manifest_rows = []
    for row in read_files():
        dest = RAW_IPEDS / f"{row['file']}.zip"
        manifest_row = download(session, row["url"], dest)
        manifest_rows.append(manifest_row | {"file_used": extract(row["file"], dest, unitids)})
    update_manifest(manifest_rows)


def extract_all():
    # Rebuilds the extracts from zips already in data/raw, without touching the network.
    unitids = read_peer_unitids()
    for row in read_files():
        extract(row["file"], RAW_IPEDS / f"{row['file']}.zip", unitids)


def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    parser = argparse.ArgumentParser()
    parser.add_argument("step", choices=["dictionaries", "data", "extract"])
    args = parser.parse_args()
    if args.step == "dictionaries":
        download_dictionaries()
    elif args.step == "data":
        download_data()
    else:
        extract_all()
    return 0


if __name__ == "__main__":
    sys.exit(main())
