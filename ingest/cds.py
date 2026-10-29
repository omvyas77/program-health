# Common Data Set files for the 18 peers, 2024-25 and 2025-26, listed in config/cds_files.csv.
# Section C1 is pre-extracted into data/manual/cds_c1_draft.csv to be checked against each
# file by hand; a value moves to data/manual/cds_c1.csv only after it has been checked by hand.

import argparse
import csv
import logging
import re
import sys
from dataclasses import dataclass

import openpyxl
import pdfplumber
from lxml import html
from openpyxl.utils import get_column_letter
from pdfminer.pdftypes import resolve1

from ingest.http import fetch, new_session
from ingest.paths import CONFIG, RAW, ROOT
from ingest.sources import download, update_manifest

log = logging.getLogger("cds")

RAW_CDS = RAW / "cds"
DRAFT = ROOT / "data" / "manual" / "cds_c1_draft.csv"
IMAGE_READS = ROOT / "data" / "manual" / "cds_c1_image_reads.csv"
STAGES = ["applicants", "admits", "enrolled"]
GENDERS = ["men", "women", "another", "unknown"]
DRAFT_COLUMNS = (
    ["unitid", "cds_year", "fall_term"]
    + [f"{stage}_total" for stage in STAGES]
    + [f"{stage}_{r}" for r in ("in_state", "out_of_state", "international") for stage in STAGES]
    + ["source_url", "page", "verified_by", "verified_on", "notes"]
)
NUMBER = re.compile(r"^\d{1,3}(,\d{3})+$|^\d+(\.0)?$")
# Column headings in C1's gender and residency tables. PDFs can split a heading over lines
# ("OUT-OF-" above "STATE"), so the patterns also match the first fragment.
HEADINGS = {
    "men": r"^(men|males?)$",
    "women": r"^(women|females?)$",
    "another": r"^another( gender)?$",
    "unknown": r"^unknown$",
    "total": r"^total$",
    "in_state": r"^in-?\s?(state)?$",
    "out_of_state": r"^out-?\s?(of)?-?\s?(state)?$",
    "international": r"^international$",
}


def read_files():
    with (CONFIG / "cds_files.csv").open(newline="") as f:
        return list(csv.DictReader(f))


def local_name(row):
    school = re.sub(r"[^a-z]+", "_", row["school"].lower()).strip("_")
    return f"{school}_{row['cds_year']}.{row['format']}"


def download_all():
    session = new_session()
    manifest_rows = []
    for row in read_files():
        if row["format"] == "none":
            log.info("%s %s: not published", row["school"], row["cds_year"])
            continue
        url = row["url"]
        if "sharepoint.com" in url:
            # SharePoint sets a cookie on the sharing link before it will serve the file.
            fetch(session, url)
            url = f"{url}&download=1"
        dest = RAW_CDS / local_name(row)
        source = f"cds/{dest.name}"
        manifest_rows.append(download(session, url, dest, source) | {"file_used": dest.name})
    update_manifest(manifest_rows)


@dataclass
class Row:
    tokens: list  # (position, text, where): x position on a PDF page, else cell index
    code: str = ""

    @property
    def label(self):
        return " ".join(t for _, t, _ in self.tokens if number(t) is None)

    @property
    def values(self):
        return [(p, number(t), w) for p, t, w in self.tokens if number(t) is not None]


def number(text):
    text = text.strip()
    return int(float(text.replace(",", ""))) if NUMBER.match(text) else None


FIELD = re.compile(
    r"^(AP_RECD|AP_ADMT|EN_TOT)_(?:(STATE|NRES|INTL|UNK)_1ST|1ST(?:_(FT|PT))?(?:_(MEN|WMN|UNK|NON_BINARY))?)_N$"
)
FIELD_STAGE = {"AP_RECD": "applicants", "AP_ADMT": "admits", "EN_TOT": "enrolled"}
FIELD_RESIDENCY = {
    "STATE": "in_state",
    "NRES": "out_of_state",
    "INTL": "international",
    "UNK": "unknown",
}
FIELD_GENDER = {"MEN": "men", "WMN": "women", "UNK": "unknown", "NON_BINARY": "another"}
# Illinois's 2024-25 workbook lays residency out as stacked blocks with no headings, but every
# row carries its template item number.
ITEM_CODES = {
    "C117": "total",
    "C120": "in_state",
    "C123": "out_of_state",
    "C126": "international",
    "C129": "unknown",
}


def decoded(value):
    value = resolve1(value)
    return value.decode("latin-1") if isinstance(value, bytes) else value


def form_fields(path):
    # The fillable CDS template names every C1 field; filled values live in the document's form,
    # which page text leaves out.
    def walk(fields, out):
        for f in fields or []:
            f = resolve1(f)
            name, value = decoded(f.get("T")), decoded(f.get("V"))
            if isinstance(name, str) and value not in (None, ""):
                out[name] = str(value).strip()
            walk(resolve1(f.get("Kids")), out)
        return out

    with pdfplumber.open(path) as pdf:
        form = resolve1(pdf.doc.catalog.get("AcroForm"))
        fields = walk(resolve1(form.get("Fields")), {}) if form else {}
        pages = [
            i + 1 for i, p in enumerate(pdf.pages) if "who applied" in (p.extract_text() or "")
        ]
    where = f"page {pages[0]}" if pages else "form field"
    found = {}
    for name, value in fields.items():
        match = FIELD.match(name)
        if not match or number(value) is None:
            continue
        stage, residency, status, gender = match.groups()
        stage = FIELD_STAGE[stage]
        if residency:
            part = f"res_{FIELD_RESIDENCY[residency]}"
        elif gender:
            part = FIELD_GENDER[gender]
        else:
            part = "res_total"
        found[(stage, (status or "all").lower(), part)] = (number(value), f"{where}, field {name}")
    return found


def pdf_rows(path):
    # Words grouped into lines by vertical position, for PDFs whose values are printed text.
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            if not re.search(r"first-year", page.extract_text() or "", re.I):
                continue
            where = f"page {i + 1}"
            tokens = []
            for w in page.extract_words():
                # Some CDS PDFs (Iowa's and Minnesota's 2024-25) use a font whose characters
                # extract as (cid:N), offset by 29 from ASCII: (cid:88)(cid:81) reads "un".
                text = re.sub(r"\(cid:(\d+)\)", lambda m: chr(int(m.group(1)) + 29), w["text"])
                previous = tokens[-1] if tokens else None
                # USC's PDF splits "3,928" into "3" and ",928"; join them back.
                if (
                    previous
                    and re.fullmatch(r",\d{3}", text)
                    and re.fullmatch(r"\d{1,3}", previous[3])
                ):
                    if abs(previous[0] - w["top"]) < 2 and w["x0"] - previous[4] < 4:
                        tokens[-1] = (
                            previous[0],
                            previous[1],
                            w["x1"],
                            previous[3] + text,
                            w["x1"],
                        )
                        continue
                tokens.append((w["top"], w["x0"], w["x1"], text, w["x1"]))
            tokens = [(top, (x0 + x1) / 2, text) for top, x0, x1, text, _ in tokens]
            line, last_top = [], None
            for top, x, text in sorted(tokens):
                if last_top is not None and top - last_top > 4:
                    yield Row(sorted(line))
                    line = []
                line.append((x, text, where))
                last_top = top
            if line:
                yield Row(sorted(line))


def xlsx_rows(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(ws for ws in workbook.worksheets if ws.title.strip() == "CDS-C")
    for r, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        # Illinois's and Wisconsin's 2025-26 workbooks repeat C1 as a separate list from column
        # AA on, misaligned against the form; only the form itself (columns A to Y) is read.
        cells = cells[:25]
        codes = [str(c).strip() for c in cells if re.fullmatch(r"C\d{3}", str(c or "").strip())]
        tokens = [
            (j, str(c).strip(), f"{sheet.title}!{get_column_letter(j + 1)}{r}")
            for j, c in enumerate(cells)
            if c not in (None, "") and not re.fullmatch(r"C\d{3}", str(c).strip())
        ]
        yield Row(tokens, codes[0] if codes else "")


def html_rows(path):
    for tr in html.fromstring(path.read_bytes()).xpath("//tr"):
        cells = [" ".join(c.text_content().split()) for c in tr.xpath("./td|./th")]
        yield Row([(j, c, "web page") for j, c in enumerate(cells) if c])


def stage_of(label):
    low = label.lower()
    if "first-time, first-year" not in low:
        return None
    for stage, pattern in (
        ("applicants", "applied"),
        ("admits", "admitted"),
        ("enrolled", "enrolled"),
    ):
        if pattern in low:
            return stage
    # Iowa's 2025-26 PDF breaks "who were / admitted" across two lines.
    if re.search(r"who\s?were$", low):
        return "admits"
    return None


def gender_of(label):
    low = label.lower()
    for gender, pattern in (
        ("men", r"\bmen\b|\bmales?\b"),
        ("women", r"\bwomen\b|\bfemales?\b"),
        ("another", r"another gender"),
        ("unknown", r"unknowns? (gender|sex)|of unknown"),
    ):
        if re.search(pattern, low):
            return gender
    return None


def headings_in(row):
    found = {}
    for pos, text, _ in row.tokens:
        for name, pattern in HEADINGS.items():
            if name not in found and re.match(pattern, text.strip().lower()):
                found[name] = pos
    return found


def read_c1(rows, tolerance):
    found = {}  # (stage, status, part) -> (value, where)
    columns = None
    block = None
    seen_c1 = False
    for k, row in enumerate(rows):
        if seen_c1 and re.match(r"C\s?2\b|C2[.\- ]", row.label.strip()):
            break
        seen_c1 = seen_c1 or bool(stage_of(row.label))
        if row.code in ITEM_CODES:
            block = ITEM_CODES[row.code]
        if block and row.code and stage_of(row.label) and len(row.values) == 1:
            part = "res_total" if block == "total" else f"res_{block}"
            found.setdefault((stage_of(row.label), "all", part), row.values[0][1:])
            continue
        stage = stage_of(row.label)
        if not row.values and not stage:
            headings = headings_in(row)
            if len(headings) >= 2:
                # Gather fragments from the lines just above and below a PDF heading.
                for near in rows[max(k - 1, 0) : k + 3]:
                    if not near.values and not stage_of(near.label):
                        for name, pos in headings_in(near).items():
                            headings.setdefault(name, pos)
                columns = headings
            continue
        if not stage or not row.values:
            continue
        low = row.label.lower()
        status = "ft" if "full-time, first" in low else "pt" if "part-time, first" in low else "all"
        gender = gender_of(row.label)
        if gender:
            value, where = row.values[-1][1:]
            found.setdefault((stage, status, gender), (value, where))
            continue
        if columns:
            for pos, value, where in row.values:
                name = min(columns, key=lambda c: abs(columns[c] - pos))
                if abs(columns[name] - pos) <= tolerance:
                    residency = any(
                        c in columns for c in ("in_state", "out_of_state", "international")
                    )
                    if residency:
                        part = f"res_{name}"
                    else:
                        part = "table_total" if name == "total" else name
                    found.setdefault((stage, status, part), (value, where))
            continue
        if len(row.values) == 1:
            found.setdefault((stage, status, "stated_total"), row.values[0][1:])
        elif len(row.values) in (4, 5) and "degree-seeking" in low:
            # A residency row whose headings did not come through. Read it in the template's
            # order (in-state, out-of-state, international, [unknown], total) only when the parts
            # add up to the last number.
            numbers = [v for _, v, _ in row.values]
            if sum(numbers[:-1]) == numbers[-1]:
                names = ["in_state", "out_of_state", "international", "unknown"][: len(numbers) - 1]
                for name, (_, value, where) in zip(names + ["total"], row.values, strict=True):
                    found.setdefault((stage, status, f"res_{name}"), (value, where))
    return found


def summarize(found):
    values, wheres, notes = {}, set(), []
    for stage in STAGES:
        candidates = {}
        for label, part in (("residency table total", "res_total"), ("table total", "table_total")):
            if (stage, "all", part) in found and found[(stage, "all", part)][0] > 0:
                candidates[label] = found[(stage, "all", part)]
        if (stage, "all", "stated_total") in found:
            candidates["stated total"] = found[(stage, "all", "stated_total")]
        genders = [found[(stage, "all", g)] for g in GENDERS if (stage, "all", g) in found]
        if (stage, "all", "men") in found and (stage, "all", "women") in found:
            candidates["sum of gender rows"] = (sum(v for v, _ in genders), genders[0][1])
        status_rows = [
            found[k] for k in found if k[0] == stage and k[1] in ("ft", "pt") and k[2] in GENDERS
        ]
        if stage == "enrolled" and status_rows:
            candidates["sum of full- and part-time rows"] = (
                sum(v for v, _ in status_rows),
                status_rows[0][1],
            )
        if candidates:
            source, (value, where) = next(iter(candidates.items()))
            values[f"{stage}_total"] = value
            wheres.add(where)
            disagree = {k: v for k, (v, _) in candidates.items() if v != value}
            if disagree:
                notes.append(f"{stage}: {source} {value:,} differs from {disagree}")
            elif source != "residency table total":
                notes.append(f"{stage} total is the {source}")
        for part in ("in_state", "out_of_state", "international"):
            if (stage, "all", f"res_{part}") in found:
                value, where = found[(stage, "all", f"res_{part}")]
                values[f"{stage}_{part}"] = value
                wheres.add(where)
        if found.get((stage, "all", "res_unknown"), (0,))[0]:
            notes.append(
                f"{stage}: {found[(stage, 'all', 'res_unknown')][0]:,} of unknown residency"
            )
    if not any(f"applicants_{p}" in values for p in ("in_state", "out_of_state")):
        notes.append("residency split not reported or not found")
    return values, sorted(wheres), notes


def extract_draft():
    rows_out = []
    for row in read_files():
        record = {c: "" for c in DRAFT_COLUMNS}
        fall = int(row["cds_year"][:4])
        record |= {
            "unitid": row["unitid"],
            "cds_year": row["cds_year"],
            "fall_term": f"Fall {fall}",
        }
        record |= {"source_url": row["url"], "notes": row["note"]}
        if row["format"] != "none":
            path = RAW_CDS / local_name(row)
            found = form_fields(path) if row["format"] == "pdf" else {}
            if len(found) < 6:
                reader = {"pdf": pdf_rows, "xlsx": xlsx_rows, "html": html_rows}[row["format"]]
                found = read_c1(list(reader(path)), tolerance=40 if row["format"] == "pdf" else 0)
            values, wheres, notes = summarize(found)
            record |= {k: f"{v}" for k, v in values.items()}
            record["page"] = "; ".join(wheres)
            record["notes"] = "; ".join(n for n in [row["note"], *notes] if n)
            log.info("%s %s: %d values", row["school"], row["cds_year"], len(values))
        rows_out.append(record)
    # Ohio State's PDFs and UCLA's 2024-25 PDF draw their numbers as shapes, so those values
    # were read from rendered page images and are listed with their page in IMAGE_READS.
    with IMAGE_READS.open(newline="") as f:
        for read in csv.DictReader(f):
            record = next(
                r
                for r in rows_out
                if (r["unitid"], r["cds_year"]) == (read["unitid"], read["cds_year"])
            )
            assert not record[read["column"]], (
                f"{read['unitid']} {read['column']} already has a value"
            )
            record[read["column"]] = read["value"]
            if read["page"] not in record["page"]:
                record["page"] = "; ".join(p for p in [record["page"], read["page"]] if p)
            if "read from the page image" not in record["notes"]:
                note = "values read from the page image, since this PDF draws its numbers as shapes"
                record["notes"] = "; ".join(n for n in [record["notes"], note] if n)
    with DRAFT.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=DRAFT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows_out)


def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    parser = argparse.ArgumentParser()
    parser.add_argument("step", choices=["download", "draft"])
    args = parser.parse_args()
    if args.step == "download":
        download_all()
    else:
        extract_draft()
    return 0


if __name__ == "__main__":
    sys.exit(main())
